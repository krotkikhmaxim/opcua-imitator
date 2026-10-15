"""Плейбэк XLSX-трендов: маппинг колонок, транспорт, взаимное исключение.

Источник проверяется на маленьком файле, собранном openpyxl: сопоставление
по двум последним сегментам пути, честный отчёт unmatched/ambiguous,
пропуск nan. Транспорт — на виртуальном времени (как тесты движка): сон
двигает часы, поэтому файл проигрывается мгновенно и детерминированно.
Для пауз/seek — сон без хода времени: плейбэк стоит на паузе ожидания,
пока тест управляет им вручную.
"""

import asyncio
import os
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

os.environ["OPC_UA_MODE"] = "imitator"

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.services.playback.service import PlaybackError, PlaybackService
from app.services.playback.xlsx import PlaybackSourceError, open_source

K1_NODE = "Chanel1.Application.GVL.GlobalInOutSignal.x10A3DigitalInput.K1Ch1_FbRtp1"
SPEED_NODE = "Chanel1.Application.GVL.DriveGVL.ProfibusData.MotorSpeed"

CONFIGS = {
    "ch1-10a3-k1": {"id": "ch1-10a3-k1", "type": "Boolean", "node_id": K1_NODE},
    "profibus-motor-speed": {"id": "profibus-motor-speed", "type": "Int16", "node_id": SPEED_NODE},
}


def _trend_path(tmp_path, rows, headers=(("[0:0]", "[0:4]", "[0:9]"),)):
    """Файл тренда: шапка-маркеры, строка путей, строки (время, k1, speed)."""
    path = tmp_path / "trend.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.worksheets[0]
    for header in headers:
        ws.append(list(header))
    ws.append(["time", K1_NODE, SPEED_NODE])
    ws.append(["mg/s", "", ""])
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    return str(path)


T0 = "23.03.2026 15:00:00.000"


def _short_rows():
    return [
        (T0, 0, 0),
        ("23.03.2026 15:00:00.300", 1, 0),
        ("23.03.2026 15:00:00.600", 1, 1500),
        ("nan", None, None),  # мусорная строка — не данные
        ("23.03.2026 15:00:00.900", "nan", 1500),  # nan в k1 — пропуск значения
        ("23.03.2026 15:00:01.200", 0, 1500),
    ]


# -- источник ----------------------------------------------------------------


def test_source_maps_by_last_two_segments_and_reports_the_rest(tmp_path):
    path = _trend_path(
        tmp_path,
        _short_rows(),
        headers=(("[0:0]", "[0:4]", "[0:9]", "[0:11]"),),
    )
    # четвёртая колонка строки путей — сигнал, которого нет в биндингах
    wb = openpyxl.load_workbook(path)
    wb.worksheets[0].cell(row=2, column=4, value="Application.GVL.Unknown.Signal")
    wb.save(path)

    src = open_source(path, CONFIGS, max_bytes=10 * 1024 * 1024)
    try:
        assert [(c.signal_id, c.signal_type) for c in src.columns] == [
            ("ch1-10a3-k1", "Boolean"),
            ("profibus-motor-speed", "Int16"),
        ]
        assert src.unmatched == ["Application.GVL.Unknown.Signal"]
    finally:
        src.close()


def test_source_streams_values_coerced_and_skips_nan(tmp_path):
    path = _trend_path(tmp_path, _short_rows())
    src = open_source(path, CONFIGS, max_bytes=10 * 1024 * 1024)
    try:
        stream = list(src.rows())
    finally:
        src.close()

    assert [at for at, _values in stream] == [0, 300, 600, 900, 1200]
    assert stream[0][1] == {"ch1-10a3-k1": False, "profibus-motor-speed": 0}
    assert stream[2][1] == {"ch1-10a3-k1": True, "profibus-motor-speed": 1500}
    assert stream[3][1] == {"profibus-motor-speed": 1500}  # nan в k1 пропущен
    assert stream[4][1] == {"ch1-10a3-k1": False, "profibus-motor-speed": 1500}
    assert src.t0 == datetime(2026, 3, 23, 15, 0, 0)


def test_source_reports_t0_at_open(tmp_path):
    path = _trend_path(tmp_path, _short_rows())
    src = open_source(path, CONFIGS, max_bytes=10 * 1024 * 1024)
    try:
        assert src.t0 == datetime(2026, 3, 23, 15, 0, 0)
    finally:
        src.close()


def test_source_rejects_missing_file_and_oversize(tmp_path):
    with pytest.raises(PlaybackSourceError) as missing:
        open_source(str(tmp_path / "nope.xlsx"), CONFIGS, max_bytes=1024)
    assert missing.value.code == "not_found"

    path = _trend_path(tmp_path, _short_rows())
    with pytest.raises(PlaybackSourceError) as big:
        open_source(path, CONFIGS, max_bytes=10)
    assert big.value.code == "too_large"


# -- транспорт ---------------------------------------------------------------


class VirtualTime:
    """Часы, которые двигает сон: файл проигрывается мгновенно."""

    def __init__(self):
        self.mono = 100.0
        self.wall = datetime(2026, 3, 23, 12, 0, 0)

    def monotonic(self):
        return self.mono

    def clock(self):
        return self.wall

    async def sleep(self, delay):
        self.mono += delay


class FrozenTime(VirtualTime):
    """Сон не двигает часы: плейбэк стоит, пока тест им управляет."""

    async def sleep(self, delay):
        await asyncio.sleep(0)


class FakeWriter:
    def __init__(self):
        self.rows = []

    async def __call__(self, signal_id, value, when):
        self.rows.append((signal_id, value, when))
        return True


def _service(time, writer):
    return PlaybackService(
        writer,
        monotonic=time.monotonic,
        sleep=time.sleep,
        clock=time.clock,
    )


async def _wait_phase(svc, phases=("finished", "error"), limit=100000):
    for _ in range(limit):
        if svc.status()["phase"] in phases:
            return svc.status()
        await asyncio.sleep(0)
    raise AssertionError(f"плейбэк не дошёл до {phases}")


def test_status_exposes_t0_of_loaded_source(tmp_path):
    async def scenario():
        path = _trend_path(tmp_path, _short_rows())
        time, writer = VirtualTime(), FakeWriter()
        svc = _service(time, writer)
        loaded = svc.load(path, CONFIGS, max_bytes=1024 * 1024)
        assert loaded["t0"] == datetime(2026, 3, 23, 15, 0, 0)
        assert svc.unload()["t0"] is None  # без источника начало неизвестно

    asyncio.run(scenario())


def test_play_restricts_to_window_between_start_and_end(tmp_path):
    """Участок «начало–конец»: до начала — скан без записей, конец включителен."""
    async def scenario():
        path = _trend_path(tmp_path, _short_rows())
        time, writer = VirtualTime(), FakeWriter()
        svc = _service(time, writer)
        svc.load(path, CONFIGS, max_bytes=1024 * 1024)
        status = svc.play(start_ms=300, end_ms=900)
        assert status["start_ms"] == 300
        assert status["end_ms"] == 900
        status = await _wait_phase(svc)
        assert status["phase"] == "finished"
        assert status["position_ms"] == 900  # строка конца написана
        assert status["duration_ms"] == 900
        assert status["rows_played"] == 4  # 2 строки скана + 2 участка
        # значения на старте участка — текущим временем (семантика seek),
        # дальше только изменения внутри участка
        assert [r[0] for r in writer.rows] == [
            "ch1-10a3-k1",
            "profibus-motor-speed",
            "profibus-motor-speed",
        ]
        assert [r[1] for r in writer.rows] == [True, 0, 1500]

    asyncio.run(scenario())


def test_play_holds_tempo_through_quiet_rows(tmp_path):
    """Тихий участок (значения не меняются) идёт темпом тренда: строки ждут
    своего момента, а не пролистываются мгновенно до конца окна. Ответ play
    при этом не показывает «позицию от нулевой базы часов»."""

    async def scenario():
        rows = [
            (T0, 0, 0),
            ("23.03.2026 15:00:00.300", 0, 0),  # тишина
            ("23.03.2026 15:00:00.600", 0, 0),  # тишина
            ("23.03.2026 15:00:00.900", 1, 0),  # единственное изменение
        ]
        path = _trend_path(tmp_path, rows)
        time, writer = VirtualTime(), FakeWriter()
        svc = _service(time, writer)
        svc.load(path, CONFIGS, max_bytes=1024 * 1024)
        status = svc.play(end_ms=900)
        assert status["position_ms"] == 0  # позиция от базы пуска, не от часов
        status = await _wait_phase(svc)
        assert status["phase"] == "finished"
        assert status["position_ms"] == 900
        # часы прошли 900 мс трендового времени, хотя менялся один сигнал
        assert time.mono == pytest.approx(100.9, abs=1e-6)

    asyncio.run(scenario())


def test_play_rejects_window_end_not_after_start(tmp_path):
    async def scenario():
        path = _trend_path(tmp_path, _short_rows())
        time, writer = VirtualTime(), FakeWriter()
        svc = _service(time, writer)
        svc.load(path, CONFIGS, max_bytes=1024 * 1024)
        with pytest.raises(PlaybackError) as bad:
            svc.play(start_ms=500, end_ms=500)
        assert bad.value.code == "bad_window"

    asyncio.run(scenario())


def test_play_writes_only_changes_with_recorded_timestamps(tmp_path):
    async def scenario():
        path = _trend_path(tmp_path, _short_rows())
        time, writer = VirtualTime(), FakeWriter()
        svc = _service(time, writer)
        svc.load(path, CONFIGS, max_bytes=1024 * 1024)
        svc.play()
        status = await _wait_phase(svc)

        assert status["phase"] == "finished"
        assert status["duration_ms"] == 1200
        # t0-кадр целиком, затем только смены; nan не создаёт записи
        assert writer.rows == [
            ("ch1-10a3-k1", False, datetime(2026, 3, 23, 15, 0, 0)),
            ("profibus-motor-speed", 0, datetime(2026, 3, 23, 15, 0, 0)),
            ("ch1-10a3-k1", True, datetime(2026, 3, 23, 15, 0, 0, 300000)),
            ("profibus-motor-speed", 1500, datetime(2026, 3, 23, 15, 0, 0, 600000)),
            ("ch1-10a3-k1", False, datetime(2026, 3, 23, 15, 0, 1, 200000)),
        ]
        assert svc.bus_held is False  # финал отпускает шину

    asyncio.run(scenario())


def test_pause_holds_position_and_resume_continues(tmp_path):
    async def scenario():
        path = _trend_path(tmp_path, _short_rows())
        time, writer = FrozenTime(), FakeWriter()
        svc = _service(time, writer)
        svc.load(path, CONFIGS, max_bytes=1024 * 1024)
        svc.play()
        await asyncio.sleep(0)
        assert svc.bus_held is True

        paused = svc.pause()
        assert paused["paused"] is True
        position_at_pause = paused["position_ms"]

        resumed = svc.resume()
        assert resumed["paused"] is False
        assert svc.position_ms() >= position_at_pause

        svc.stop()
        assert svc.bus_held is False

    asyncio.run(scenario())


def test_seek_writes_target_state_at_wall_time_and_continues_by_trend(tmp_path):
    async def scenario():
        # 11 строк по 100 мс: t=0 и t=100..1000
        rows = [(T0, 0, 0)]
        for i in range(1, 11):
            ms = i * 100
            stamp = f"23.03.2026 15:00:00.{ms:03d}" if ms < 1000 else (
                f"23.03.2026 15:00:01.{ms % 1000:03d}"
            )
            rows.append((stamp, i % 2, i * 10))

        path = _trend_path(tmp_path, rows)
        time, writer = FrozenTime(), FakeWriter()
        svc = _service(time, writer)
        svc.load(path, CONFIGS, max_bytes=1024 * 1024)
        svc.play()
        await asyncio.sleep(0)

        status = await svc.seek(700)
        assert status["position_ms"] == 700
        # новая задача досканивает до цели и пишет значения текущим временем
        for _ in range(10000):
            if any(r[2] == time.wall for r in writer.rows):
                break
            await asyncio.sleep(0)
        flushed = [r for r in writer.rows if r[2] == time.wall]
        # значения на цели — текущим wall-временем; скорость на 700 мс равна 70
        assert ("profibus-motor-speed", 70, time.wall) in flushed

        svc.stop()

    asyncio.run(scenario())


def test_play_twice_needs_stop_and_replays_from_zero(tmp_path):
    async def scenario():
        path = _trend_path(tmp_path, _short_rows())
        time, writer = VirtualTime(), FakeWriter()
        svc = _service(time, writer)
        svc.load(path, CONFIGS, max_bytes=1024 * 1024)
        svc.play()
        with pytest.raises(PlaybackError) as busy:
            svc.play()
        assert busy.value.code == "already_running"
        await _wait_phase(svc)

        svc.play()  # финал отпустил шину — перезапуск с нуля допустим
        status = await _wait_phase(svc)
        assert status["writes"] == 5  # оба прохода по 5 записей
        assert writer.rows[-1][2] == datetime(2026, 3, 23, 15, 0, 1, 200000)

    asyncio.run(scenario())


def test_rate_is_validated_and_reported(tmp_path):
    async def scenario():
        path = _trend_path(tmp_path, _short_rows())
        time, writer = FrozenTime(), FakeWriter()
        svc = _service(time, writer)
        svc.load(path, CONFIGS, max_bytes=1024 * 1024)
        with pytest.raises(PlaybackError) as idle:
            svc.set_rate(2.0)
        assert idle.value.code == "not_running"

        svc.play()
        await asyncio.sleep(0)
        with pytest.raises(PlaybackError) as bad:
            svc.set_rate(16.0)
        assert bad.value.code == "bad_rate"
        assert svc.set_rate(4.0)["rate"] == 4.0
        svc.stop()

    asyncio.run(scenario())


def test_control_on_empty_service_is_rejected():
    async def scenario():
        time, writer = VirtualTime(), FakeWriter()
        svc = _service(time, writer)
        with pytest.raises(PlaybackError) as play:
            svc.play()
        assert play.value.code == "not_loaded"
        with pytest.raises(PlaybackError) as seek:
            await svc.seek(0)
        assert seek.value.code == "not_running"

    asyncio.run(scenario())


# -- HTTP-слой ---------------------------------------------------------------


@pytest.fixture()
def client():
    from app.main import app

    with TestClient(app) as test_client:
        yield test_client


def test_http_load_play_blocks_manual_write_then_stop_releases(tmp_path, client, monkeypatch):
    """Взаимное исключение «плейбэк ↔ ручная запись» при включённом плейбэке.

    По умолчанию (ТЗ #5, §2) плейбэк заморожен — это пинуется отдельно в
    test_playback_freeze.py; здесь флаг singleton-сервиса на время теста
    снимается, чтобы механика блокировок осталась покрытой. Объект сервиса
    не подменяется: предикат «playback» в реестре связан с самим объектом.
    """
    from app import main as app_main

    monkeypatch.setattr(app_main._playback, "_writes_enabled", True)

    path = _trend_path(tmp_path, _short_rows())

    loaded = client.post("/api/playback/load_path", json={"path": path})
    assert loaded.status_code == 200
    body = loaded.json()
    assert body["mapped_signals"] == 2
    assert body["source"] == "trend.xlsx"

    missing = client.post("/api/playback/load_path", json={"path": str(tmp_path / "no.xlsx")})
    assert missing.status_code == 404

    played = client.post("/api/playback/control", json={"action": "play"})
    assert played.status_code == 200

    conflict = client.post("/api/signals/write", json={"id": "ch1-10a3-k1", "value": False})
    assert conflict.status_code == 409
    assert "playback" in conflict.json()["detail"]

    stopped = client.post("/api/playback/control", json={"action": "stop"})
    assert stopped.status_code == 200
    assert client.post("/api/playback/unload").status_code == 200

    allowed = client.post("/api/signals/write", json={"id": "ch1-10a3-k1", "value": False})
    assert allowed.status_code == 200
